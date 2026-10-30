import os
import tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.models.marketplace import IdentitySubmission, OwnerApplication
from app.models.user import User
from app.services.accounts import display_name
STATUS_FILL = {
    "approved": PatternFill("solid", fgColor="D8EFE3"),
    "rejected": PatternFill("solid", fgColor="F6DADA"),
    "submitted": PatternFill("solid", fgColor="FBF0D9"),
    "pending": PatternFill("solid", fgColor="FBF0D9"),
}
STATUS_LABEL = {"submitted": "Pending", "pending": "Pending", "approved": "Verified", "rejected": "Rejected"}
def _fmt(dt):
    return dt.strftime("%Y-%m-%d %H:%M UTC") if dt else ""
def _sheet(ws, headers, rows, status_col):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D5B4C")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
        cell = ws.cell(row=ws.max_row, column=status_col)
        key = next((k for k, v in STATUS_LABEL.items() if v == cell.value), None)
        if key:
            cell.fill = STATUS_FILL[key]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 48)
async def rebuild_register(db: AsyncSession) -> Path:
    users = {u.id: u for u in await db.scalars(select(User))}
    owners = list(await db.scalars(select(OwnerApplication).order_by(OwnerApplication.submitted_at.desc())))
    ids = list(await db.scalars(select(IdentitySubmission).order_by(IdentitySubmission.submitted_at.desc())))
    wb = Workbook()
    ws = wb.active
    ws.title = "Owner verification"
    rows = []
    for a in owners:
        u = users.get(a.user_id)
        rows.append([
            _fmt(a.submitted_at), display_name(u) if u else "", u.username if u else "", u.email if u else "",
            u.phone if u else "", a.property_address, a.relation, a.id_doc_type, a.property_doc_type,
            STATUS_LABEL.get(a.status, a.status), _fmt(a.reviewed_at), _reviewer(users, a.reviewed_by),
            a.review_note or "", a.id_file_hash, a.property_file_hash, a.user_id,
        ])
    _sheet(ws, ["Submitted", "Name", "Username", "Email", "Phone", "Property address", "Relation", "ID document",
                "Property document", "Status", "Reviewed", "Reviewed by", "Note", "ID doc fingerprint",
                "Property doc fingerprint", "User ID"], rows, status_col=10)
    ws2 = wb.create_sheet("ID verification")
    rows = []
    for s in ids:
        u = users.get(s.user_id)
        rows.append([
            _fmt(s.submitted_at), display_name(u) if u else "", u.username if u else "", u.email if u else "",
            u.phone if u else "", s.doc_type, STATUS_LABEL.get(s.status, s.status), _fmt(s.reviewed_at),
            _reviewer(users, s.reviewed_by), s.review_note or "", s.file_hash, s.user_id,
        ])
    _sheet(ws2, ["Submitted", "Name", "Username", "Email", "Phone", "Document", "Status", "Reviewed",
                 "Reviewed by", "Note", "Document fingerprint", "User ID"], rows, status_col=7)
    target = Path(settings.VERIFICATION_REGISTER_PATH)
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
    os.close(fd)
    wb.save(tmp)
    os.replace(tmp, target)
    return target
def _reviewer(users, reviewed_by):
    if not reviewed_by:
        return ""
    u = users.get(reviewed_by)
    return display_name(u) if u else reviewed_by
async def try_rebuild_register(db: AsyncSession) -> None:
    try:
        await rebuild_register(db)
    except PermissionError:
        pass
