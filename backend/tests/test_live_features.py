import hashlib
import hmac
import io
import uuid
import pytest
from openpyxl import load_workbook
from sqlalchemy import update
from app.core.config import settings
from app.core.database import AsyncSessionLocal, normalise_database_url
from app.models.deposits import Deposit
from app.services import payments
from app.services.seed import seed_demo_admin
from tests.test_marketplace import owner_with_listing, signup, slot
async def admin_headers(client):
    async with AsyncSessionLocal() as db:
        await seed_demo_admin(db)
    res = await client.post("/api/v1/auth/password-login", json={"identifier": "admin", "password": settings.DEMO_ADMIN_PASSWORD})
    assert res.status_code == 200, res.text
    assert res.json()["user"]["isAdmin"] is True
    return {"Authorization": f"Bearer {res.json()['access_token']}"}
async def test_admin_desk_and_excel_register(client, monkeypatch):
    ah = await admin_headers(client)
    user, uh = await signup(client)
    monkeypatch.setattr(settings, "DEMO_MODE", False)
    res = await client.post("/api/v1/users/me/identity", headers=uh, data={"docType": "PAN card"},
                            files={"file": ("pan.jpg", io.BytesIO(b"pan-bytes"), "image/jpeg")})
    assert res.json()["pendingReview"] is True and res.json()["user"]["idVerified"] is False
    blocked = await client.post("/api/v1/payments/order", headers=uh, json={"purpose": "premium", "planId": "week"})
    assert blocked.status_code == 403
    assert (await client.get("/api/v1/admin/verifications", headers=uh)).status_code == 403
    desk = (await client.get("/api/v1/admin/verifications", headers=ah)).json()
    sub = next(s for s in desk["identities"] if s["userId"] == user["id"])
    assert sub["status"] == "pending" and sub["applicant"]["email"] == user["email"]
    doc = await client.get(f"/api/v1/admin/identity/{sub['id']}/file", headers=ah)
    assert doc.content == b"pan-bytes" and doc.headers["content-type"] == "image/jpeg"
    ok = await client.post(f"/api/v1/admin/identity/{sub['id']}/review", headers=ah, json={"approve": True})
    assert ok.json()["submission"]["status"] == "approved"
    assert (await client.get("/api/v1/users/me", headers=uh)).json()["user"]["idVerified"] is True
    xlsx = await client.get("/api/v1/admin/verification-register.xlsx", headers=ah)
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Owner verification", "ID verification"]
    rows = list(wb["ID verification"].iter_rows(values_only=True))
    assert rows[0][3] == "Email"
    mine = next(r for r in rows[1:] if r[3] == user["email"])
    assert mine[6] == "Verified" and mine[8] == "TrustRent Admin"
def test_razorpay_signature_and_db_url(monkeypatch):
    monkeypatch.setattr(settings, "RAZORPAY_KEY_SECRET", "secret123")
    sig = hmac.new(b"secret123", b"order_1|pay_1", hashlib.sha256).hexdigest()
    assert payments.verify_signature("order_1", "pay_1", sig)
    assert not payments.verify_signature("order_1", "pay_2", sig)
    url, args = normalise_database_url("postgres://u:p@db.example.com/trustrent?sslmode=require")
    assert url == "postgresql+asyncpg://u:p@db.example.com/trustrent" and args == {"ssl": True}
async def test_signup_email_code_over_smtp(client, monkeypatch):
    from tests.test_marketplace import profile
    sent = {}
    async def fake_send(to, code, minutes):
        sent["to"], sent["code"] = to, code
    monkeypatch.setattr(settings, "SMTP_USER", "trustrent.otp@gmail.com")
    monkeypatch.setattr(settings, "SMTP_APP_PASSWORD", "abcdabcdabcdabcd")
    monkeypatch.setattr("app.api.v1.accounts.send_code_email", fake_send)
    t = uuid.uuid4().hex[:6]
    form = profile(t) | {"email": f"Riya_{t}@Example.com"}
    res = await client.post("/api/v1/auth/signup/email-code", json=form)
    assert res.status_code == 200 and "demoCode" not in res.json()
    assert sent["to"] == f"riya_{t}@example.com" and len(sent["code"]) == 6 and sent["code"].isdigit()
    assert (await client.post("/api/v1/auth/signup/email-code", json=form)).status_code == 429
    ok = await client.post("/api/v1/auth/signup", json=form | {"emailCode": sent["code"]})
    assert ok.status_code == 201 and ok.json()["user"]["email"] == f"riya_{t}@example.com"
async def test_signup_email_failure_and_no_smtp(client, monkeypatch):
    from app.services.mailer import EmailError, build_code_email
    from tests.test_marketplace import profile
    async def broken(to, code, minutes):
        raise EmailError("We couldn't send the email. Try again in a minute.")
    monkeypatch.setattr(settings, "SMTP_USER", "trustrent.otp@gmail.com")
    monkeypatch.setattr(settings, "SMTP_APP_PASSWORD", "abcdabcdabcdabcd")
    monkeypatch.setattr("app.api.v1.accounts.send_code_email", broken)
    form = profile(uuid.uuid4().hex[:6])
    res = await client.post("/api/v1/auth/signup/email-code", json=form)
    assert res.status_code == 502 and "couldn't send" in res.json()["detail"]["message"]
    assert (await client.post("/api/v1/auth/signup/email-code", json=form)).status_code == 502
    msg = build_code_email("a@b.com", "123456", 10)
    assert msg["Subject"] == "123456 is your TrustRent verification code"
    monkeypatch.setattr(settings, "SMTP_APP_PASSWORD", "")
    off = await client.post("/api/v1/auth/signup/email-code", json=profile(uuid.uuid4().hex[:6]))
    assert off.status_code == 503
async def verified_tenant_with_visit(client, listing, oh):
    tenant, th = await signup(client)
    await client.post("/api/v1/users/me/identity", headers=th, data={"docType": "Passport"},
                      files={"file": ("p.jpg", io.BytesIO(b"passport"), "image/jpeg")})
    order = await client.post("/api/v1/payments/order", headers=th, json={"purpose": "premium", "planId": "week"})
    await client.post("/api/v1/payments/verify", headers=th, json={"orderId": order.json()["orderId"]})
    early = await client.post("/api/v1/payments/order", headers=th, json={"purpose": "deposit", "listingId": listing["id"]})
    assert early.status_code == 403
    conv = (await client.post("/api/v1/conversations", headers=th, json={"listingId": listing["id"]})).json()
    visit = (await client.post(f"/api/v1/conversations/{conv['id']}/messages", headers=th,
                               json={"kind": "visit", "slots": [slot(2)]})).json()
    await client.post(f"/api/v1/messages/{visit['id']}/visit-response", headers=oh,
                      json={"accept": True, "slot": visit["visit"]["slots"][0]})
    return tenant, th
async def test_deposit_receipt_lifecycle(client):
    owner, oh, listing = await owner_with_listing(client)
    tenant, th = await verified_tenant_with_visit(client, listing, oh)
    order = await client.post("/api/v1/payments/order", headers=th, json={"purpose": "deposit", "listingId": listing["id"]})
    assert order.status_code == 200 and order.json()["amount"] == 34000 and order.json()["provider"] == "demo"
    paid = await client.post("/api/v1/payments/verify", headers=th, json={"orderId": order.json()["orderId"]})
    dep_id = paid.json()["depositId"]
    assert dep_id.startswith("DEP-")
    again = await client.post("/api/v1/payments/order", headers=th, json={"purpose": "deposit", "listingId": listing["id"]})
    assert again.status_code == 409
    dep = (await client.get(f"/api/v1/deposits/{dep_id}", headers=th)).json()
    assert dep["status"] == "held" and dep["receipt"]["network"] == "trustrent-ledger"
    assert dep["verification"]["verified"] is True
    tx = dep["receipt"]["recordTx"]
    by_hash = (await client.get("/api/v1/deposits/lookup", headers=oh, params={"q": tx})).json()
    assert by_hash["id"] == dep_id and by_hash["tenant"]["name"] == "Riya Sharma"
    stranger, sh = await signup(client)
    public = (await client.get("/api/v1/deposits/lookup", headers=sh, params={"q": dep_id})).json()
    assert "tenant" not in public and public["amount"] == 34000
    bad = await client.post(f"/api/v1/deposits/{dep_id}/propose", headers=oh, json={"deducted": 40000, "reason": "too much"})
    assert bad.status_code == 400
    prop = await client.post(f"/api/v1/deposits/{dep_id}/propose", headers=oh,
                             json={"deducted": 4000, "reason": "Broken bathroom tiles and two months of unpaid electricity"})
    assert prop.json()["status"] == "settlement_proposed" and prop.json()["refunded"] == 30000
    disputed = await client.post(f"/api/v1/deposits/{dep_id}/respond", headers=th, json={"accept": False})
    assert disputed.json()["status"] == "disputed"
    ah = await admin_headers(client)
    settled = await client.post(f"/api/v1/admin/deposits/{dep_id}/settle", headers=ah,
                                json={"deducted": 2000, "reason": "Tiles only; electricity was already paid"})
    s = settled.json()
    assert s["status"] == "settled" and s["refunded"] == 32000 and s["ownerPayout"] == "due"
    assert s["receipt"]["settleTx"] and s["verification"]["verified"] is True
    paid_out = await client.post(f"/api/v1/admin/deposits/{dep_id}/owner-paid", headers=ah)
    assert paid_out.json()["ownerPayout"] == "paid"
    async with AsyncSessionLocal() as db:
        await db.execute(update(Deposit).where(Deposit.id == dep_id).values(refunded=34000, deducted=0))
        await db.commit()
    check = (await client.get("/api/v1/deposits/lookup", headers=th, params={"q": dep_id})).json()
    assert check["verification"]["verified"] is False and check["verification"]["ledgerIntact"] is True
async def test_tenant_accepts_settlement(client):
    owner, oh, listing = await owner_with_listing(client)
    tenant, th = await verified_tenant_with_visit(client, listing, oh)
    order = await client.post("/api/v1/payments/order", headers=th, json={"purpose": "deposit", "listingId": listing["id"]})
    dep_id = (await client.post("/api/v1/payments/verify", headers=th, json={"orderId": order.json()["orderId"]})).json()["depositId"]
    await client.post(f"/api/v1/deposits/{dep_id}/propose", headers=oh, json={"deducted": 0, "reason": ""})
    done = (await client.post(f"/api/v1/deposits/{dep_id}/respond", headers=th, json={"accept": True})).json()
    assert done["status"] == "settled" and done["refunded"] == 34000 and done["ownerPayout"] is None
    mine = (await client.get("/api/v1/deposits", headers=oh)).json()
    assert any(d["id"] == dep_id and d["role"] == "owner" for d in mine)
@pytest.mark.parametrize("demo", [False])
async def test_payments_off_without_keys(client, monkeypatch, demo):
    _, uh = await signup(client)
    await client.post("/api/v1/users/me/identity", headers=uh, data={"docType": "PAN card"},
                      files={"file": ("pan.jpg", io.BytesIO(b"x"), "image/jpeg")})
    monkeypatch.setattr(settings, "DEMO_MODE", demo)
    res = await client.post("/api/v1/payments/order", headers=uh, json={"purpose": "premium", "planId": "month"})
    assert res.status_code == 503
    cfg = (await client.get("/api/v1/config")).json()
    assert cfg["payments"] == "off" and cfg["depositNetwork"] == "trustrent-ledger"
